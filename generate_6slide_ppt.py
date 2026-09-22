import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

WORKDIR = os.path.dirname(os.path.abspath(__file__))

def create_6slide_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    NAVY = RGBColor(30, 58, 138)
    DARK = RGBColor(15, 23, 42)
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(248, 250, 252)
    BORDER_COLOR = RGBColor(203, 213, 225)
    GREEN = RGBColor(21, 128, 61)

    img_home = os.path.join(WORKDIR, "slide_img_home.png")
    img_modal = os.path.join(WORKDIR, "slide_img_modal.png")
    img_joined = os.path.join(WORKDIR, "slide_img_joined.png")
    img_card = os.path.join(WORKDIR, "slide_img_card.png")

    slides_config = [
        # SLIDE 1
        {
            "title": "Campus Club Explorer",
            "is_title_slide": True,
            "subtitle": "A React-Based Web Application for College Club Discovery & Management",
            "meta_items": [
                "Department of Computer Science & Engineering",
                "College Lab Training Assessment Project",
                "Built with React 18, Vite & Solid-Color CSS3",
                "Features: Props • Lists & Keys • Filters • Conditional Rendering"
            ],
            "image": img_home,
            "notes": "Respected Sir and faculty members, this is my React training assessment project titled 'Campus Club Explorer'. It is a single-page web portal designed for college students to discover clubs, filter by domain, search by keywords, view meeting details, and manage club memberships."
        },
        # SLIDE 2
        {
            "title": "Problem Statement & Solution Architecture",
            "is_title_slide": False,
            "bullets": [
                ("Campus Problem", "Club activities and meeting schedules are fragmented across noticeboards and chat groups, leading to low student awareness and participation."),
                ("Project Solution", "A centralized, interactive web application allowing students to browse clubs across Coding, Sports, Arts, and Entrepreneurship."),
                ("Technology Stack", "React 18 for component architecture, Vite for sub-second hot reloading, and solid-color CSS3 for clean aesthetics."),
                ("Unidirectional Architecture", "App.jsx serves as the central state hub, passing data down via props and receiving user events via callback functions.")
            ],
            "image": img_card,
            "image_caption": "Live Dashboard with Category Filters & Club Cards",
            "notes": "Sir, on our campus students often miss club meetings because info is scattered. Campus Club Explorer centralizes all information. Architecturally, App.jsx manages state and passes data down to child components following React's unidirectional flow."
        },
        # SLIDE 3
        {
            "title": "React Concepts 1 & 2: Props, Reuse, Lists & Keys",
            "is_title_slide": False,
            "bullets": [
                ("Props (Data Passing)", "Props are read-only inputs passed from App.jsx to ClubCard (club, isJoined, onViewDetails). This decouples UI from state."),
                ("Component Reuse", "The ClubCard component is authored once and dynamically reused across all 8 campus clubs with zero duplicate markup."),
                ("Lists Rendering (.map)", "JavaScript's .map() loops through the filtered clubs array and returns a <ClubCard /> for every matched item."),
                ("The Unique 'key' Prop", "Each item receives key={club.id}. This enables React's Virtual DOM reconciliation engine to track item updates efficiently.")
            ],
            "image": img_card,
            "image_caption": "Component Reuse: 8 Clubs Rendered via .map() with unique keys",
            "notes": "Sir, here are two major concepts: Props and Lists. Instead of writing 8 separate cards, we created one reusable ClubCard component. We map over the array using club.id as the key, which lets React re-render only the exact card that changes."
        },
        # SLIDE 4
        {
            "title": "React Concept 3: Dynamic Search & Category Filtering",
            "is_title_slide": False,
            "bullets": [
                ("Controlled Form Inputs", "searchQuery and selectedCategory are maintained in React state using the useState hook."),
                ("Simultaneous Dual-Predicate Filter", "The filtering function evaluates category matching ('All' or exact) AND keyword matching (.includes()) together on every keystroke."),
                ("Zero Page Reloads", "Filtered results update reactively in real time without refreshing the browser."),
                ("User Ergonomics", "Includes an instant 'Clear' search button and active category pill buttons with visual state highlights.")
            ],
            "image": img_home,
            "image_caption": "Search Bar & Category Pills (Coding, Sports, Arts, E-Cell)",
            "notes": "Our search and filter UI evaluates both the active category button and the search text simultaneously. When a student selects 'Coding' and types 'AI', it filters instantly without page reloads using clean array methods."
        },
        # SLIDE 5
        {
            "title": "React Concept 4: Conditional Rendering & Modal Dialog",
            "is_title_slide": False,
            "bullets": [
                ("Short-Circuit Modal Rendering", "{selectedClub && <ClubModal ... />} ensures zero modal DOM overhead until a club is clicked."),
                ("Empty State Fallback", "Ternary operator renders the club grid when items exist, or a friendly 'No clubs found' box with a 'Reset Filters' button."),
                ("Comprehensive Club View", "Modal displays meeting schedule, faculty coordinator advisor, member count, and bulleted activities list."),
                ("Accessible Interaction", "Modal closes gracefully on clicking the 'Close' button, 'Done' button, or tapping the darkened background.")
            ],
            "image": img_modal,
            "image_caption": "Interactive Club Details Modal with Advisor & Meeting Schedule",
            "notes": "Sir, we demonstrated Conditional Rendering in multiple places: showing the modal only when selectedClub is not null, showing an empty state message if no clubs match, and closing the modal cleanly on background click."
        },
        # SLIDE 6
        {
            "title": "Interactive Membership, Test Verification & Conclusion",
            "is_title_slide": False,
            "bullets": [
                ("Interactive Join / Leave", "Clicking 'Join Club' toggles membership, increments member count, and changes button to 'Leave Club'."),
                ("Persistent 'JOINED' Tag", "Card immediately displays a green 'JOINED' badge on the dashboard, tracking student membership."),
                ("Real-Time Popup Alerts", "Displays green toast alert: 'Success: You have joined [Club]!' auto-dismissed in 3 seconds."),
                ("Test Verification & Build", "All 8 test cases PASS. Production build completes with zero errors in under 500ms via Vite.")
            ],
            "image": img_joined,
            "image_caption": "Live Joined Status Banner, Leave Button & Real-Time Feedback",
            "notes": "Finally, we implemented interactive membership tracking. When you join, it shows a popup message, updates member count, and marks the card with a JOINED badge. All test cases passed with zero build errors. Thank you, Sir!"
        }
    ]

    for idx, sdata in enumerate(slides_config):
        slide = prs.slides.add_slide(blank_layout)

        # Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT_BG
        bg.line.fill.background()

        # Speaker notes
        slide.notes_slide.notes_text_frame.text = sdata.get("notes", "")

        if sdata.get("is_title_slide"):
            # Header block
            hbar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.3))
            hbar.fill.solid()
            hbar.fill.fore_color.rgb = NAVY
            hbar.line.fill.background()

            tbox = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.733), Inches(0.9))
            tf = tbox.text_frame
            p = tf.paragraphs[0]
            p.text = "CAMPUS CLUB EXPLORER"
            p.font.size = Pt(36)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.font.name = "Arial"

            # Left content box
            lbox = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.4))
            lbox.fill.solid()
            lbox.fill.fore_color.rgb = WHITE
            lbox.line.color.rgb = BORDER_COLOR
            lbox.line.width = Pt(1.5)

            ltf = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.4), Inches(5.0)).text_frame
            ltf.word_wrap = True

            sp1 = ltf.paragraphs[0]
            sp1.text = "React Web Application Project"
            sp1.font.size = Pt(20)
            sp1.font.bold = True
            sp1.font.color.rgb = NAVY

            sp2 = ltf.add_paragraph()
            sp2.text = sdata["subtitle"]
            sp2.font.size = Pt(14)
            sp2.font.color.rgb = DARK
            sp2.space_before = Pt(10)

            sp3 = ltf.add_paragraph()
            sp3.text = "Project Highlights & Assessment Info:"
            sp3.font.size = Pt(14)
            sp3.font.bold = True
            sp3.font.color.rgb = NAVY
            sp3.space_before = Pt(16)

            for m in sdata["meta_items"]:
                mp = ltf.add_paragraph()
                mp.text = f"• {m}"
                mp.font.size = Pt(12)
                mp.font.color.rgb = DARK
                mp.space_before = Pt(6)

            # Right image box
            if os.path.exists(sdata["image"]):
                # Container card for image
                icard = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.4))
                icard.fill.solid()
                icard.fill.fore_color.rgb = WHITE
                icard.line.color.rgb = BORDER_COLOR
                icard.line.width = Pt(1.5)

                # Add picture inside
                slide.shapes.add_picture(sdata["image"], Inches(6.95), Inches(1.9), Inches(5.4), Inches(4.7))

        else:
            # Top Banner
            top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
            top_bar.fill.solid()
            top_bar.fill.fore_color.rgb = NAVY
            top_bar.line.fill.background()

            # Slide Number indicator
            snum_box = slide.shapes.add_textbox(Inches(11.5), Inches(0.3), Inches(1.2), Inches(0.5))
            snp = snum_box.text_frame.paragraphs[0]
            snp.text = f"Slide {idx + 1} / 6"
            snp.font.size = Pt(14)
            snp.font.bold = True
            snp.font.color.rgb = RGBColor(191, 219, 254)

            # Slide Title
            tbox = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(10.5), Inches(0.75))
            ttf = tbox.text_frame
            tp = ttf.paragraphs[0]
            tp.text = sdata["title"]
            tp.font.size = Pt(24)
            tp.font.bold = True
            tp.font.color.rgb = WHITE
            tp.font.name = "Arial"

            # Left side: Detailed Content Card
            left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(6.4), Inches(5.65))
            left_card.fill.solid()
            left_card.fill.fore_color.rgb = WHITE
            left_card.line.color.rgb = BORDER_COLOR
            left_card.line.width = Pt(1.5)

            ctext = slide.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(6.0), Inches(5.35))
            ctf = ctext.text_frame
            ctf.word_wrap = True

            for bidx, (bhead, bdesc) in enumerate(sdata.get("bullets", [])):
                hp = ctf.paragraphs[0] if bidx == 0 else ctf.add_paragraph()
                hp.text = f"{bidx + 1}. {bhead}"
                hp.font.size = Pt(13.5)
                hp.font.bold = True
                hp.font.color.rgb = NAVY
                if bidx > 0:
                    hp.space_before = Pt(10)

                dp = ctf.add_paragraph()
                dp.text = bdesc
                dp.font.size = Pt(11.5)
                dp.font.color.rgb = DARK
                dp.space_before = Pt(2)

            # Right side: Image Card
            right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.4), Inches(1.4), Inches(5.1), Inches(5.65))
            right_card.fill.solid()
            right_card.fill.fore_color.rgb = WHITE
            right_card.line.color.rgb = BORDER_COLOR
            right_card.line.width = Pt(1.5)

            if os.path.exists(sdata.get("image", "")):
                slide.shapes.add_picture(sdata["image"], Inches(7.55), Inches(1.6), Inches(4.8), Inches(4.65))

            # Image caption
            cap_box = slide.shapes.add_textbox(Inches(7.4), Inches(6.35), Inches(5.1), Inches(0.6))
            cap_p = cap_box.text_frame.paragraphs[0]
            cap_p.text = sdata.get("image_caption", "Actual UI Screenshot")
            cap_p.font.size = Pt(10.5)
            cap_p.font.bold = True
            cap_p.font.color.rgb = NAVY
            cap_p.alignment = PP_ALIGN.CENTER

    out_file = os.path.join(WORKDIR, "Campus_Club_Explorer_6Slide_Presentation.pptx")
    prs.save(out_file)
    print(f"6-Slide PPT generated successfully at: {out_file}")
    return out_file

if __name__ == "__main__":
    create_6slide_presentation()
